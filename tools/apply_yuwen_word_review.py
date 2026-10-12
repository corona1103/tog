import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path('/Users/tal/Desktop/tog')
SQL_DIR = ROOT / 'sql'
REVIEW_XLSX = ROOT / 'yuwen' / '语文积累本内容清单.xlsx'

VERSION_IDS = {
    '部编版': '00b6bf17df4345109f570bf0273705f3',
    '统编版（2024）六三制': '5e21a87b47a648eb9da56a51381e1611',
}


def extract_tuples(text):
    out = []
    i = 0
    while True:
        i = text.find('\n(', i)
        if i < 0:
            return out
        start = i + 1
        j = start
        quoted = False
        while j < len(text):
            char = text[j]
            if char == "'":
                if quoted and j + 1 < len(text) and text[j + 1] == "'":
                    j += 2
                    continue
                quoted = not quoted
            elif not quoted and char == ')':
                out.append(text[start:j + 1])
                i = j + 1
                break
            j += 1
        else:
            return out


def split_tuple(value):
    body = value[1:-1]
    fields = []
    start = 0
    quoted = False
    i = 0
    while i < len(body):
        char = body[i]
        if char == "'":
            if quoted and i + 1 < len(body) and body[i + 1] == "'":
                i += 2
                continue
            quoted = not quoted
        elif not quoted and char == ',':
            fields.append(body[start:i].strip())
            start = i + 1
        i += 1
    fields.append(body[start:].strip())
    return fields


def sql_value(value):
    if value == 'NULL':
        return None
    if value.startswith("'") and value.endswith("'"):
        return value[1:-1].replace("''", "'").replace('\\\\', '\\')
    try:
        return int(value)
    except ValueError:
        return value


def read_rows(filename, field_count):
    text = (SQL_DIR / filename).read_text()
    return [
        [sql_value(field) for field in split_tuple(item)]
        for item in extract_tuples(text)
        if len(split_tuple(item)) == field_count and re.fullmatch(r'\d+', split_tuple(item)[0])
    ]


def sql_string(value):
    if value is None:
        return 'NULL'
    if isinstance(value, bool):
        return '1' if value else '0'
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace('\\', '\\\\').replace("'", "''").replace('\x00', '') + "'"


def render_insert(table, columns, rows, header, pre_insert=None, post_insert=None):
    lines = [header, '', 'START TRANSACTION;']
    if pre_insert:
        lines.extend(['', *pre_insert])
    lines.extend([
        '',
        f'INSERT INTO `{table}`',
        '(' + ', '.join(f'`{column}`' for column in columns) + ')',
        'VALUES',
        ',\n'.join('(' + ', '.join(sql_string(value) for value in row) + ')' for row in rows),
        'ON DUPLICATE KEY UPDATE',
    ])
    lines.extend(',\n'.join(f'`{column}` = VALUES(`{column}`)' for column in columns[1:]).split('\n'))
    lines[-1] += ';'
    if post_insert:
        lines.extend(['', *post_insert])
    lines.extend(['', 'COMMIT;', ''])
    return '\n'.join(lines)


def parse_labels(text):
    result = {}
    if not text:
        return result
    matches = list(re.finditer(r'(?m)(原词|释义|篇目|来源|原文|译文):\s*', str(text)))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(str(text))
        result[match.group(1)] = str(text)[match.end():end].strip().strip('/').strip()
    return result


def clean_meaning(value):
    return str(value or '').strip().rstrip('。')


def clean_title(value):
    value = str(value or '').strip()
    value = value.replace('《', '').replace('》', '').replace('<', '').replace('>', '')
    return re.sub(r'\s+', '', value)


def normalize_source(value, fallback_title=''):
    parts = [part.strip() for part in str(value or '').split('｜') if part.strip()]
    title_match = re.search(r'《([^》]+)》', str(value or ''))
    title = clean_title(title_match.group(1) if title_match else fallback_title)
    volume = parts[1] if len(parts) >= 4 else (parts[0] if parts else '')
    return f'{volume}｜《{title}》'


def mark_word(text, word):
    text = str(text or '').strip()
    if not text or not word:
        return text
    return text.replace(word, f'<b>{word}</b>')


def normalize_er_meaning(value):
    value = clean_meaning(value)
    for label in ('并列', '递进', '转折', '修饰', '承接', '因果'):
        if f'表{label}' in value:
            return f'<连>表{label}'
    return value


def parse_option_updates(text):
    updates = {}
    for line in str(text or '').splitlines():
        line = line.strip()
        match = re.match(r'^([A-D])(?:选项)?(?:换成|改成)?[：:]\s*(.+)$', line)
        if not match:
            match = re.match(r'^([A-D])[\.．]\s*(.+)$', line)
        if match:
            updates[match.group(1)] = match.group(2).strip()
    return updates


def parse_example_updates(text, word):
    result = {}
    for line in str(text or '').splitlines():
        match = re.match(r'^例句([12])[：:]\s*(.*)$', line.strip())
        if match:
            result[int(match.group(1))] = mark_word(match.group(2).strip(), word)
    return result


def plain(value):
    return re.sub(r'<[^>]+>', '', str(value or ''))


def main():
    sense_rows = read_rows('edu_chn_word_sense.sql', 10)
    question_rows = read_rows('edu_chn_word_question.sql', 7)
    catalog_rows = read_rows('edu_chn_word_catalog.sql', 9)
    sense_by_id = {row[0]: row for row in sense_rows}

    workbook = load_workbook(REVIEW_XLSX, read_only=True, data_only=True)
    sheet = workbook['实虚词一审']
    reviews = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        raw = str(row[1] or '')
        id_match = re.search(r'edu_chn_word_sense\.id: (\d+)', raw)
        correction = str(row[2] or '')
        fields = parse_labels(correction)
        reviews.append({
            'id': int(id_match.group(1)) if id_match else None,
            'word': row[0],
            'correction': correction,
            'fields': fields,
            'q1': str(row[3] or ''),
            'q2': str(row[4] or ''),
        })

    deleted_ids = set()
    changed_ids = set()
    new_rows = []
    review_by_id = {}
    max_sense_id = max(row[0] for row in sense_rows)

    for review in reviews:
        sense_id = review['id']
        fields = review['fields']
        if sense_id is not None:
            if '删去' in review['correction']:
                deleted_ids.add(sense_id)
                continue
            if sense_id not in sense_by_id:
                continue
            target = sense_by_id[sense_id]
            word = target[2]
            if '释义' in fields:
                target[4] = clean_meaning(fields['释义'])
            if '原文' in fields:
                target[5] = mark_word(fields['原文'], word)
                changed_ids.add(sense_id)
            if '译文' in fields:
                target[7] = fields['译文']
            if '释义' in fields or '译文' in fields:
                changed_ids.add(sense_id)
            review['target_id'] = sense_id
            review_by_id[sense_id] = review
            continue

        if '新增' not in review['correction']:
            continue
        word = clean_title(fields.get('原词') or review.get('word'))
        title = clean_title(fields.get('篇目'))
        source = normalize_source(fields.get('来源'), title)
        meaning = clean_meaning(fields.get('释义'))
        if word == '而':
            word_class = 2
        else:
            candidates = [row for row in sense_rows if row[2] == word]
            word_class = candidates[0][1] if candidates else 1
        if word == '而':
            meaning = normalize_er_meaning(meaning)
        example = mark_word(fields.get('原文'), word)
        trans = fields.get('译文', '')
        duplicate = next((row for row in sense_rows + new_rows if row[2] == word and row[4] == meaning and row[5] == example), None)
        if duplicate:
            review['target_id'] = duplicate[0]
            review_by_id[duplicate[0]] = review
            continue
        max_sense_id += 1
        target = [max_sense_id, word_class, word, next((row[3] for row in sense_rows if row[2] == word), ''), meaning, example, source, trans, 0, 0]
        new_rows.append(target)
        review['target_id'] = max_sense_id
        review_by_id[max_sense_id] = review

    active_senses = [row for row in sense_rows if row[0] not in deleted_ids] + new_rows
    for row in active_senses:
        if row[2] == '而':
            row[4] = normalize_er_meaning(row[4])

    # 1085 的审核备注中把“译文”标签误带进了释义字段，这里按最终确认内容
    # 固定拆回正确的释义与译文，避免后续重新生成时再次混入题目选项。
    for row in active_senses:
        if row[0] == 1085:
            row[4] = '<动>隔绝；隔断'
            row[7] = '下行、上行的航道都被阻断。'

    # “而”按用户确认的高中《劝学》范围固定保留六类常考用法。
    # Excel 修订表中未单独落一条“表递进”，这里补齐该词义，避免被现有
    # 其他来源的“而”记录覆盖或遗漏。
    er_progressive = '<连>表递进'
    er_progressive_example = '君子博学<b>而</b>日参省乎己，则知明而行无过矣。'
    er_progressive_source = '高中必修上册｜《劝学》'
    er_progressive_trans = '君子广泛学习并且每天检查反省自己，就能智慧明达，行为没有过错。'
    er_progressive_review = {
        'id': None,
        'fields': {'来源': er_progressive_source},
        'q1': 'A: <连>表递进\nB: <连>表并列\nC: <连>表转折\nD: <连>表修饰',
        'q2': '例句1：暂无例句\n例句2：青，取之于蓝，而青于蓝；',
    }
    if not any(
        row[2] == '而' and row[4] == er_progressive and row[5] == er_progressive_example
        for row in active_senses
    ):
        max_sense_id += 1
        er_progressive_row = [
            max_sense_id, 2, '而', 'ér', er_progressive,
            er_progressive_example, er_progressive_source,
            er_progressive_trans, 0, 0,
        ]
        new_rows.append(er_progressive_row)
        active_senses.append(er_progressive_row)
    for row in active_senses:
        if row[2] == '而' and row[4] == er_progressive:
            review_by_id[row[0]] = er_progressive_review

    # 题目按词义修订表覆盖；题型一首项固定为当前正确释义，题型二首项固定为正确例句。
    questions = {(row[1], row[2]): row for row in question_rows if row[1] not in deleted_ids}
    max_question_id = max((row[0] for row in question_rows), default=0)
    active_by_id = {row[0]: row for row in active_senses}
    for sense in active_senses:
        sense_id, _, word, _, meaning, example, _, _, _, _ = sense
        review = review_by_id.get(sense_id)
        if not example:
            questions.pop((sense_id, 1), None)
            questions.pop((sense_id, 2), None)
            continue
        q1 = questions.get((sense_id, 1))
        q2 = questions.get((sense_id, 2))
        if q1 is None:
            max_question_id += 1
            q1 = [max_question_id, sense_id, 1, '', '', 1, 0]
            questions[(sense_id, 1)] = q1
        if q2 is None:
            max_question_id += 1
            q2 = [max_question_id, sense_id, 2, '', '', 1, 0]
            questions[(sense_id, 2)] = q2

        q1_options = json.loads(q1[4]) if q1[4] else []
        if not q1_options:
            q1_options = [{'content': meaning, 'correct': True}]
        q1_options[0] = {'content': meaning, 'correct': True}
        while len(q1_options) < 4:
            q1_options.append({'content': '暂无释义', 'correct': False})
        q1_options = q1_options[:4]
        for option in q1_options[1:]:
            option['correct'] = False

        q2_options = json.loads(q2[4]) if q2[4] else []
        if not q2_options:
            q2_options = [
                {'content': '暂无例句', 'correct': True},
                {'content': '暂无例句', 'correct': False},
            ]
        while len(q2_options) < 2:
            q2_options.append({'content': '暂无例句', 'correct': False})
        q2_options = q2_options[:2]
        q2_options[0]['correct'] = True
        q2_options[1]['correct'] = False

        if review:
            option_updates = parse_option_updates(review['q1'])
            for label, value in option_updates.items():
                index = ord(label) - ord('A')
                if index == 0:
                    q1_options[0]['content'] = meaning
                elif index < len(q1_options):
                    q1_options[index]['content'] = value
            example_updates = parse_example_updates(review['q2'], word)
            if example_updates:
                # 审核表是增量修订：只填写了哪一侧，就只替换哪一侧；
                # 未填写的一侧沿用原题目内容，不能误改成“暂无例句”。
                for index, content in example_updates.items():
                    if index <= len(q2_options):
                        q2_options[index - 1]['content'] = content
                    else:
                        q2_options.append({'content': content, 'correct': index == 1})
            if '原文' in review['fields'] or review['id'] is None:
                q1[3] = f'下列对例句“{plain(example)}”中“{word}”的解释，正确的一项是'

        q1[3] = q1[3] or f'下列对例句“{plain(example)}”中“{word}”的解释，正确的一项是'
        q2[3] = q2[3] or f'下列句子中“{word}”的意义和用法，与例句相同的一项是'
        q1[4] = json.dumps(q1_options, ensure_ascii=False, separators=(',', ':'))
        q2[4] = json.dumps(q2_options, ensure_ascii=False, separators=(',', ':'))

    active_questions = sorted(questions.values(), key=lambda row: row[0])

    catalog_by_sense = {row[5]: row for row in catalog_rows}
    volume_map = {(row[2], row[4]): (row[1], row[3]) for row in catalog_rows}
    max_catalog_id = max(row[0] for row in catalog_rows)
    max_sort = {}
    for row in catalog_rows:
        key = (row[1], row[3])
        max_sort[key] = max(max_sort.get(key, 0), row[6])
    new_catalogs = []
    for sense in new_rows:
        review = review_by_id[sense[0]]
        source = review['fields'].get('来源', '')
        parts = [part.strip() for part in source.split('｜')]
        version_name = parts[0] if len(parts) >= 4 else '部编版'
        volume_name = parts[1] if len(parts) >= 4 else source.split('｜')[0]
        if version_name == '部编版' and not volume_name.startswith('高中'):
            version_name = '统编版（2024）六三制'
        version_id, volume_id = volume_map.get((version_name, volume_name), (VERSION_IDS.get(version_name, ''), ''))
        if not volume_id:
            raise ValueError(f'新增词义无法匹配教材分册：{source}')
        key = (version_id, volume_id)
        max_sort[key] = max_sort.get(key, 0) + 1
        max_catalog_id += 1
        new_catalogs.append([max_catalog_id, version_id, version_name, volume_id, volume_name, sense[0], max_sort[key], 1, 0])

    active_catalogs = [row for row in catalog_rows if row[5] not in deleted_ids] + new_catalogs

    deleted_sense_ids = ', '.join(str(value) for value in sorted(deleted_ids))
    sense_pre = [f'UPDATE `edu_chn_word_sense` SET `status` = 0, `is_del` = 1 WHERE `id` IN ({deleted_sense_ids});'] if deleted_ids else None
    question_pre = [f'UPDATE `edu_chn_word_question` SET `status` = 0, `is_del` = 1 WHERE `sense_id` IN ({deleted_sense_ids});'] if deleted_ids else None
    catalog_pre = [f'UPDATE `edu_chn_word_catalog` SET `status` = 0, `is_del` = 1 WHERE `sense_id` IN ({deleted_sense_ids});'] if deleted_ids else None

    sense_sql = render_insert(
        'edu_chn_word_sense',
        ['id', 'word_class', 'word', 'pinyin', 'meaning', 'example', 'source', 'example_trans', 'status', 'is_del'],
        sorted(active_senses, key=lambda row: row[0]),
        '-- 语文积累本-高中与初中实虚词词义例句（统一数据表，按实虚词一审修订）',
        sense_pre,
    )
    question_sql = render_insert(
        'edu_chn_word_question',
        ['id', 'sense_id', 'question_type', 'stem', 'options_json', 'status', 'is_del'],
        active_questions,
        '-- 语文积累本-高中与初中实虚词题目（按实虚词一审修订）',
        question_pre,
    )
    catalog_prefix = (SQL_DIR / 'edu_chn_word_catalog.sql').read_text().split('START TRANSACTION;')[0].rstrip()
    catalog_sql = catalog_prefix + '\n\n' + render_insert(
        'edu_chn_word_catalog',
        ['id', 'version_id', 'version_name', 'volume_id', 'volume_name', 'sense_id', 'sort', 'status', 'is_del'],
        sorted(active_catalogs, key=lambda row: row[0]),
        '',
        catalog_pre,
    )

    (SQL_DIR / 'edu_chn_word_sense.sql').write_text(sense_sql)
    (SQL_DIR / 'edu_chn_word_question.sql').write_text(question_sql)
    (SQL_DIR / 'edu_chn_word_catalog.sql').write_text(catalog_sql)
    print(json.dumps({
        'reviews': len(reviews),
        'deleted_senses': len(deleted_ids),
        'new_senses': len(new_rows),
        'active_senses': len(active_senses),
        'active_questions': len(active_questions),
        'active_catalogs': len(active_catalogs),
        'er_meanings': sorted({row[4] for row in active_senses if row[2] == '而'}),
    }, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
