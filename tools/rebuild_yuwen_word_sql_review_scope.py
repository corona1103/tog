"""按《审核问题修改.xlsx》的增量范围重建实虚词 SQL。

基线固定取 HEAD 中三张 SQL 的有效状态，当前工作区 SQL 只作为审核覆盖层：
已有 sense_id 只在审核表出现时覆盖；工作区新增的 sense_id 必须是本次审核新增项。
"""

from __future__ import annotations

import re
import subprocess
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SQL_DIR = ROOT / "sql"
REVIEW_XLSX = ROOT / "yuwen" / "审核问题修改.xlsx"


def split_top(value: str) -> list[str]:
    result: list[str] = []
    start = 0
    quoted = False
    index = 0
    while index < len(value):
        char = value[index]
        if char == "'":
            if quoted and index + 1 < len(value) and value[index + 1] == "'":
                index += 2
                continue
            quoted = not quoted
        elif not quoted and char == ",":
            result.append(value[start:index].strip())
            start = index + 1
        index += 1
    result.append(value[start:].strip())
    return result


def sql_value(value: str):
    value = value.strip()
    if value == "NULL":
        return None
    if value.startswith("'") and value.endswith("'"):
        return value[1:-1].replace("''", "'").replace("\\\\", "\\")
    try:
        return int(value)
    except ValueError:
        return value


def statement_end(text: str, start: int) -> int:
    quoted = False
    index = start
    while index < len(text):
        char = text[index]
        if char == "'":
            if quoted and index + 1 < len(text) and text[index + 1] == "'":
                index += 2
                continue
            quoted = not quoted
        elif not quoted and char == ";":
            return index
        index += 1
    return len(text)


def parse_tuples(body: str, field_count: int) -> list[list]:
    rows: list[list] = []
    index = 0
    while index < len(body):
        if body[index] != "(":
            index += 1
            continue
        start = index
        index += 1
        quoted = False
        while index < len(body):
            char = body[index]
            if char == "'":
                if quoted and index + 1 < len(body) and body[index + 1] == "'":
                    index += 2
                    continue
                quoted = not quoted
            elif not quoted and char == ")":
                fields = split_top(body[start + 1:index])
                if len(fields) == field_count:
                    rows.append([sql_value(field) for field in fields])
                index += 1
                break
            index += 1
        else:
            break
    return rows


def parse_sql_state(text: str, table: str, field_count: int) -> dict[int, dict]:
    """按 SQL 语句顺序还原 INSERT/UPDATE 的最终状态。"""
    state: dict[int, dict] = {}
    events = []
    pattern = re.compile(rf"(?m)^(INSERT INTO|UPDATE)\s+`?{re.escape(table)}`?")
    for match in pattern.finditer(text):
        end = statement_end(text, match.start())
        events.append((match.start(), match.group(1), text[match.start():end + 1]))

    for _, kind, statement in sorted(events):
        if kind == "INSERT INTO":
            columns_match = re.search(
                rf"INSERT INTO\s+`?{re.escape(table)}`?\s*\((.*?)\)\s*VALUES\s*",
                statement,
                re.S,
            )
            if not columns_match:
                continue
            columns = [item.strip().strip("`") for item in split_top(columns_match.group(1))]
            values_start = columns_match.end()
            duplicate_start = statement.find("\nON DUPLICATE KEY UPDATE", values_start)
            values_end = duplicate_start if duplicate_start >= 0 else len(statement) - 1
            rows = parse_tuples(statement[values_start:values_end], field_count)

            duplicate_fields: dict[str, object] = {}
            if duplicate_start >= 0:
                duplicate_part = statement[duplicate_start + len("\nON DUPLICATE KEY UPDATE") : -1]
                for assignment in split_top(duplicate_part):
                    match = re.match(r"`?([A-Za-z_]+)`?\s*=\s*(.+)$", assignment.strip(), re.S)
                    if not match:
                        continue
                    field, expression = match.groups()
                    if expression.strip().startswith("VALUES("):
                        duplicate_fields[field] = ("VALUES", field)
                    else:
                        duplicate_fields[field] = sql_value(expression)

            for row in rows:
                values = dict(zip(columns, row))
                key = int(values["id"])
                if key not in state:
                    state[key] = values
                    continue
                if not duplicate_fields:
                    state[key].update(values)
                    continue
                for field, expression in duplicate_fields.items():
                    if isinstance(expression, tuple) and expression[0] == "VALUES":
                        state[key][field] = values[expression[1]]
                    else:
                        state[key][field] = expression
            continue

        set_match = re.search(r"\bSET\s+", statement, re.S)
        where_match = re.search(r"\nWHERE\s+", statement, re.S)
        if not set_match or not where_match:
            continue
        set_part = statement[set_match.end() : where_match.start()]
        where_part = statement[where_match.end() :]

        ids: set[int] | None = None
        in_match = re.search(r"`?id`?\s+IN\s*\(([^)]*)\)", where_part, re.I | re.S)
        if in_match:
            ids = {int(item) for item in re.findall(r"\d+", in_match.group(1))}
        equal_match = re.search(r"`?id`?\s*=\s*(\d+)", where_part, re.I)
        if equal_match:
            ids = {int(equal_match.group(1))}
        if ids is None and re.search(r"`?is_del`?\s*=\s*0", where_part, re.I):
            ids = {key for key, row in state.items() if row.get("is_del") == 0}
        if ids is None:
            continue

        case_updates: dict[str, dict[int, object]] = {}
        for case_match in re.finditer(
            r"`?([A-Za-z_]+)`?\s*=\s*CASE\s+`?id`?([\s\S]*?)END",
            set_part,
            re.I,
        ):
            field = case_match.group(1)
            case_updates[field] = {
                int(key): sql_value(value)
                for key, value in re.findall(
                    r"WHEN\s+(\d+)\s+THEN\s+((?:'(?:''|[^'])*')|NULL|-?\d+)",
                    case_match.group(2),
                    re.I,
                )
            }

        direct_part = re.sub(
            r"`?([A-Za-z_]+)`?\s*=\s*CASE\s+`?id`?[\s\S]*?END",
            "",
            set_part,
            flags=re.I,
        )
        direct_updates = {}
        for assignment in split_top(direct_part):
            match = re.match(r"`?([A-Za-z_]+)`?\s*=\s*(.+)$", assignment.strip(), re.S)
            if match:
                direct_updates[match.group(1)] = sql_value(match.group(2).strip())

        for key in ids:
            state.setdefault(key, {"id": key})
            for field, values in case_updates.items():
                if key in values:
                    state[key][field] = values[key]
            state[key].update(direct_updates)
    return state


def git_head(path: str) -> str:
    return subprocess.check_output(["git", "show", f"HEAD:{path}"], cwd=ROOT, text=True)


def labels(text: str) -> dict[str, str]:
    result: dict[str, str] = {}
    matches = list(re.finditer(r"(?m)(原词|释义|篇目|来源|原文|译文):\s*", text or ""))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        result[match.group(1)] = text[match.end() : end].strip().strip("/").strip()
    return result


def sql_string(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    return "'" + str(value).replace("\\", "\\\\").replace("'", "''").replace("\x00", "") + "'"


def render_insert(table: str, columns: list[str], rows: list[dict], header: str, deletes: list[int]) -> str:
    lines = [header, "", "START TRANSACTION;", ""]
    if deletes:
        lines.append(
            f"UPDATE `{table}` SET `status` = 0, `is_del` = 1 WHERE `sense_id` IN ({', '.join(map(str, deletes))});"
            if table == "edu_chn_word_question"
            else f"UPDATE `{table}` SET `status` = 0, `is_del` = 1 WHERE `id` IN ({', '.join(map(str, deletes))});"
        )
        if table == "edu_chn_word_sense":
            lines[-1] = f"UPDATE `{table}` SET `status` = 0, `is_del` = 1 WHERE `id` IN ({', '.join(map(str, deletes))});"
        if table == "edu_chn_word_catalog":
            lines[-1] = f"UPDATE `{table}` SET `status` = 0, `is_del` = 1 WHERE `sense_id` IN ({', '.join(map(str, deletes))});"
        lines.append("")
    lines.extend(
        [
            f"INSERT INTO `{table}`",
            "(" + ", ".join(f"`{column}`" for column in columns) + ")",
            "VALUES",
            ",\n".join(
                "(" + ", ".join(sql_string(row.get(column)) for column in columns) + ")" for row in rows
            ),
            "ON DUPLICATE KEY UPDATE",
            ",\n".join(f"`{column}` = VALUES(`{column}`)" for column in columns[1:]) + ";",
            "",
            "COMMIT;",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    table_specs = {
        "sense": ("edu_chn_word_sense", 10, "sql/edu_chn_word_sense.sql"),
        "question": ("edu_chn_word_question", 7, "sql/edu_chn_word_question.sql"),
        "catalog": ("edu_chn_word_catalog", 9, "sql/edu_chn_word_catalog.sql"),
    }
    baseline = {}
    current = {}
    for key, (table, count, path) in table_specs.items():
        baseline[key] = parse_sql_state(git_head(path), table, count)
        current[key] = parse_sql_state((ROOT / path).read_text(), table, count)

    workbook = load_workbook(REVIEW_XLSX, read_only=True, data_only=True)
    sheet = workbook.active
    review_ids: set[int] = set()
    deleted_ids: set[int] = set()
    valid_new_rows = 0
    review_count = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        review_count += 1
        raw = str(row[1] or "")
        correction = str(row[2] or "")
        match = re.search(r"edu_chn_word_sense\.id:\s*(\d+)", raw)
        if match:
            sense_id = int(match.group(1))
            review_ids.add(sense_id)
            if "删去" in correction:
                deleted_ids.add(sense_id)
            continue
        fields = labels(correction)
        if all(fields.get(name) for name in ("原词", "释义", "篇目", "原文")):
            valid_new_rows += 1

    current_only_senses = set(current["sense"]) - set(baseline["sense"])
    unexpected_new = sorted(current_only_senses - {153} - set(range(1844, 2284)))
    if unexpected_new:
        raise RuntimeError(f"工作区存在无法归入审核新增项的词义 ID：{unexpected_new}")

    affected_senses = (review_ids - deleted_ids) | current_only_senses
    final_senses = dict(baseline["sense"])
    for sense_id in deleted_ids:
        final_senses.pop(sense_id, None)
    for sense_id in affected_senses:
        if sense_id in current["sense"] and sense_id not in deleted_ids:
            final_senses[sense_id] = current["sense"][sense_id]

    final_sense_ids = set(final_senses)

    current_questions_by_sense: dict[int, list[dict]] = defaultdict(list)
    for row in current["question"].values():
        current_questions_by_sense[row["sense_id"]].append(row)
    final_questions: dict[int, dict] = {}
    for question_id, row in baseline["question"].items():
        sense_id = row.get("sense_id")
        if sense_id not in final_sense_ids or sense_id in deleted_ids:
            continue
        if sense_id in affected_senses and current_questions_by_sense.get(sense_id):
            continue
        final_questions[question_id] = row
    for row in current["question"].values():
        sense_id = row.get("sense_id")
        if sense_id in affected_senses and sense_id in final_sense_ids:
            final_questions[row["id"]] = row

    current_catalog_by_sense: dict[int, list[dict]] = defaultdict(list)
    for row in current["catalog"].values():
        current_catalog_by_sense[row["sense_id"]].append(row)
    final_catalog: dict[int, dict] = {}
    for catalog_id, row in baseline["catalog"].items():
        sense_id = row.get("sense_id")
        if sense_id not in final_sense_ids or sense_id in deleted_ids:
            continue
        if sense_id in affected_senses and current_catalog_by_sense.get(sense_id):
            continue
        final_catalog[catalog_id] = row
    for row in current["catalog"].values():
        sense_id = row.get("sense_id")
        if sense_id in affected_senses and sense_id in final_sense_ids:
            final_catalog[row["id"]] = row

    sense_columns = ["id", "word_class", "word", "pinyin", "meaning", "example", "source", "example_trans", "status", "is_del"]
    question_columns = ["id", "sense_id", "question_type", "stem", "options_json", "status", "is_del"]
    catalog_columns = ["id", "version_id", "version_name", "volume_id", "volume_name", "sense_id", "sort", "status", "is_del"]

    (SQL_DIR / "edu_chn_word_sense.sql").write_text(
        render_insert(
            "edu_chn_word_sense",
            sense_columns,
            [final_senses[key] for key in sorted(final_senses)],
            "-- 语文积累本-高中与初中实虚词词义例句（审核表范围覆盖版）",
            sorted(deleted_ids),
        )
    )
    (SQL_DIR / "edu_chn_word_question.sql").write_text(
        render_insert(
            "edu_chn_word_question",
            question_columns,
            [final_questions[key] for key in sorted(final_questions)],
            "-- 语文积累本-高中与初中实虚词题目（审核表范围覆盖版）",
            sorted(deleted_ids),
        )
    )
    (SQL_DIR / "edu_chn_word_catalog.sql").write_text(
        render_insert(
            "edu_chn_word_catalog",
            catalog_columns,
            [final_catalog[key] for key in sorted(final_catalog)],
            "-- 语文积累本-高中与初中实虚词教材关系（审核表范围覆盖版）",
            sorted(deleted_ids),
        )
    )

    print({
        "reviews": review_count,
        "review_sense_ids": len(review_ids),
        "deleted_sense_ids": len(deleted_ids),
        "valid_new_rows": valid_new_rows,
        "current_only_senses": len(current_only_senses),
        "final_senses": len(final_senses),
        "final_questions": len(final_questions),
        "final_catalog": len(final_catalog),
    })


if __name__ == "__main__":
    main()
